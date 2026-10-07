from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
GUI_FILE = ROOT / "vega_gui.py"
OUTPUT_FILE = ROOT / "Procedure_Outil_Migration_VEGA6.docx"


def extract_metadata() -> tuple[str, str]:
    content = GUI_FILE.read_text(encoding="utf-8")
    title_match = re.search(r'APP_TITLE = "([^"]+)"', content)
    version_match = re.search(r'APP_VERSION = "([^"]+)"', content)
    title = title_match.group(1) if title_match else "Outil de migration VEGA6"
    version = version_match.group(1) if version_match else "Version inconnue"
    return title, version


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "8")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "808080")
        borders.append(border)
    tbl_pr.append(borders)


def set_row_height(row, height_cm: float) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_height = OxmlElement("w:trHeight")
    tr_height.set(qn("w:val"), str(int(height_cm * 567)))
    tr_height.set(qn("w:hRule"), "exact")
    tr_pr.append(tr_height)


def add_title(document: Document, title: str, version: str) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = "Tahoma"

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f"Procédure d'utilisation - version {version}")
    run.italic = True
    run.font.size = Pt(11)
    run.font.name = "Tahoma"

    date_line = document.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_line.add_run(f"Document généré le {date.today().strftime('%d/%m/%Y')}")
    run.font.size = Pt(10)
    run.font.name = "Tahoma"


def add_intro(document: Document) -> None:
    document.add_paragraph(
        "Ce document explique le fonctionnement général de l'outil et sert de support de prise en main "
        "pour les utilisateurs internes.",
        style="Body Text",
    )
    document.add_paragraph(
        "Les cadres grisés présents dans le document sont prévus pour coller vos captures d'écran.",
        style="Body Text",
    )


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_screenshot_box(document: Document, label: str, height_cm: float = 6.0) -> None:
    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(f"Capture à ajouter : {label}")
    run.bold = True
    run.font.name = "Tahoma"
    run.font.size = Pt(9)

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    cell = table.cell(0, 0)
    cell.width = Cm(16.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F2F2F2")
    set_row_height(table.rows[0], height_cm)

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"[ Insérer ici la capture d'écran ]\n{label}")
    run.bold = True
    run.font.name = "Tahoma"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(90, 90, 90)

    document.add_paragraph("")


def add_section(document: Document, title: str, intro: str | None = None) -> None:
    document.add_heading(title, level=1)
    if intro:
        document.add_paragraph(intro, style="Body Text")


def build_document() -> Path:
    title, version = extract_metadata()
    document = Document()

    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Tahoma"
    normal.font.size = Pt(10)

    body = document.styles["Body Text"]
    body.font.name = "Tahoma"
    body.font.size = Pt(10)

    add_title(document, title, version)
    add_intro(document)

    add_section(document, "1. Prérequis")
    add_bullets(
        document,
        [
            "Lancer l'outil en mode administrateur.",
            "Vérifier que le dossier Vega ciblé est bien la racine de l'environnement à traiter.",
            "Préparer les éléments nécessaires avant migration : arrêt HFSQL et réinitialisation Retail Force si demandé.",
        ],
    )
    add_screenshot_box(document, "Page de connexion ou ouverture de l'outil")

    add_section(document, "2. Connexion à l'outil")
    add_bullets(
        document,
        [
            "Saisir le mot de passe dans l'écran d'accès protégé.",
            "La saisie peut se faire au clavier, à la souris ou en tactile selon le poste.",
            "Une fois le mot de passe valide, l'outil s'ouvre automatiquement.",
        ],
    )
    add_screenshot_box(document, "Écran de connexion")

    add_section(document, "3. Vue générale")
    add_bullets(
        document,
        [
            "La fenêtre principale comporte plusieurs onglets : Migration, Impressions Vega, Pare-feu et Téléchargement.",
            "Le journal d'activité affiche les actions réalisées par l'outil.",
            "Le bandeau inférieur affiche l'état courant de l'application.",
        ],
    )
    add_screenshot_box(document, "Fenêtre principale")

    add_section(document, "4. Onglet Migration")
    add_bullets(
        document,
        [
            "Renseigner ou vérifier le chemin de la racine Vega.",
            "Cliquer sur Vérifier pour contrôler la structure vega.dos\\V6 et la situation du port TCP 7678.",
            "Cocher les prérequis utilisateur avant de lancer l'opération.",
            "Cliquer sur Lancer la migration lorsque le dossier est validé.",
            "En cas de besoin, utiliser la liste des sauvegardes pour lancer un retour arrière.",
        ],
    )
    add_screenshot_box(document, "Onglet Migration - vérification et lancement")
    add_screenshot_box(document, "Onglet Migration - retour arrière", height_cm=5.0)

    add_section(document, "5. Onglet Impressions Vega")
    add_bullets(
        document,
        [
            "Utiliser Vérifier pour contrôler le dossier C:\\ImpressionsVega et son partage.",
            "Utiliser Créer / réparer pour remettre en conformité le dossier et le partage si nécessaire.",
            "Contrôler les statuts affichés avant de quitter l'onglet.",
        ],
    )
    add_screenshot_box(document, "Onglet Impressions Vega")

    add_section(document, "6. Onglet Pare-feu")
    add_bullets(
        document,
        [
            "Cet onglet permet de vérifier ou créer la règle Service Fiscal RetailForce.",
            "La règle attendue concerne le protocole TCP sur le port local 7678.",
            "Utiliser Créer / réparer si la règle est absente ou mal configurée.",
        ],
    )
    add_screenshot_box(document, "Onglet Pare-feu")

    add_section(document, "7. Onglet Téléchargement")
    add_bullets(
        document,
        [
            "Sélectionner les fichiers souhaités dans la liste.",
            "Le dossier de destination par défaut correspond au dossier Téléchargements de la session Windows.",
            "Cliquer sur Télécharger les fichiers cochés pour lancer la file.",
            "Suivre la progression grâce à la barre, au pourcentage, au temps restant et à la vitesse de téléchargement.",
        ],
    )
    add_screenshot_box(document, "Onglet Téléchargement - sélection des fichiers")
    add_screenshot_box(document, "Onglet Téléchargement - progression", height_cm=5.0)

    add_section(document, "8. Informations et support")
    add_bullets(
        document,
        [
            "Le bouton Infos ouvre la page de support de l'outil.",
            "Cette page rappelle l'éditeur, la version et les moyens de contact.",
            "Les boutons permettent d'ouvrir Microsoft Teams ou de préparer un mail vers le support.",
        ],
    )
    add_screenshot_box(document, "Onglet Informations")

    add_section(document, "9. Conseils d'utilisation")
    add_bullets(
        document,
        [
            "Toujours lire le journal d'activité en cas d'avertissement ou d'erreur.",
            "Ne fermer l'outil qu'une fois l'opération terminée.",
            "Conserver les captures d'écran utiles dans ce document pour faciliter le support interne.",
        ],
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"{title} - Procédure d'utilisation - version {version}")
    footer_run.font.name = "Tahoma"
    footer_run.font.size = Pt(8)

    document.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    path = build_document()
    print(path)
